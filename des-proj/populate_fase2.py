import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = pptx.Presentation('des-proj/Template - Fase 2 - Disciplina Integradora de Desenvolvimento de Projetos.pptx')

# Slide 1: Cover
s1 = prs.slides[0]
if len(s1.shapes) > 0 and s1.shapes[0].has_text_frame:
    tf = s1.shapes[0].text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = 'Samuel Aiedo'
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

# Slide 3: Seção 1 - GitHub
s3 = prs.slides[2]
if len(s3.shapes) > 0 and s3.shapes[0].has_text_frame:
    tf = s3.shapes[0].text_frame
    tf.word_wrap = True
    tf.clear()
    
    lines = [
        ('ORGANIZAÇÃO DO REPOSITÓRIO E BRANCH DE ENTREGA DA FASE 2:', True, Pt(12), RGBColor(30, 58, 138)),
        ('• Repositório Oficial no GitHub: https://github.com/Faiedo/medqr-clinical-system', False, Pt(10), RGBColor(71, 85, 105)),
        ('• Branch de Entrega da Fase 2: https://github.com/Faiedo/medqr-clinical-system/tree/fase-2', False, Pt(10), RGBColor(71, 85, 105)),
        ('• Visibilidade: Pública (acesso irrestrito para avaliação da tutoria e docentes sem necessidade de login)', False, Pt(9.5), RGBColor(71, 85, 105)),
        ('', False, Pt(4), RGBColor(0,0,0)),
        ('MÓDULOS DE DESENVOLVIMENTO INTEGRADOS NO REPOSITÓRIO:', True, Pt(11), RGBColor(15, 23, 42)),
        ('1. Módulo Backend (Node.js & Express): API RESTful completa contendo regras de validação médica (devops/src/index.js), geração de tokens opacos de emergência e endpoint de observabilidade /health com uptime e verificação de integridade.', False, Pt(9.5), RGBColor(51, 65, 85)),
        ('2. Módulo Frontend (Aplicação Web Interativa): Interface responsiva Single-Page Application (des-proj/app/index.html) com 4 telas funcionais: Login de Trabalhador, Formulário Clínico de 9 Campos Vitais, Geração de Crachá com QR Code dinâmico e Portal Mobile do Socorrista com bloqueio por Senha Pública.', False, Pt(9.5), RGBColor(51, 65, 85)),
        ('3. Módulo de Banco de Dados (PostgreSQL 15): Scripts DDL de inicialização (devops/init-db.sql) com tabelas normalizadas usuarios e fichas_clinicas, chaves estrangeiras com integridade referencial ON DELETE CASCADE e índices de alta performance para busca rápida em emergências.', False, Pt(9.5), RGBColor(51, 65, 85)),
        ('4. Infraestrutura e Automação DevOps: Dockerfile multi-stage com usuário não-root, orquestração docker-compose.yml e pipeline de CI/CD automatizado no GitHub Actions.', False, Pt(9.5), RGBColor(51, 65, 85))
    ]
    for idx, (txt, bold, size, color) in enumerate(lines):
        p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
        p.text = txt
        p.font.bold = bold
        p.font.size = size
        p.font.color.rgb = color

# Slide 5: Seção 2 - User Stories
s5 = prs.slides[4]
if len(s5.shapes) > 0 and s5.shapes[0].has_text_frame:
    tf = s5.shapes[0].text_frame
    tf.word_wrap = True
    tf.clear()
    
    lines = [
        ('ATUALIZAÇÃO E REFINAMENTO DAS USER STORIES (FASE 1 -> FASE 2):', True, Pt(11), RGBColor(30, 58, 138)),
        ('• US01 - Cadastro de Trabalhador: Refinada com criptografia irreversível de senhas via bcrypt (salt rounds = 10) e validação de e-mail corporativo único. [Status: Concluído]', False, Pt(8.5), RGBColor(51, 65, 85)),
        ('• US02 - Autenticação Segura: Implementada sessão stateless via JSON Web Token (JWT) com verificação de assinatura e expiração controlada. [Status: Concluído]', False, Pt(8.5), RGBColor(51, 65, 85)),
        ('• US03 - Preenchimento da Ficha Clínica: Adicionada sanitização de campos e validação de 9 campos vitais no backend com retorno HTTP 400 amigável. [Status: Concluído]', False, Pt(8.5), RGBColor(51, 65, 85)),
        ('• US04 - Seleção de Tipo Sanguíneo: Criada regra de enumeração médica estrita (A+, A-, B+, B-, AB+, AB-, O+, O-), bloqueando dados inválidos. [Status: Concluído]', False, Pt(8.5), RGBColor(51, 65, 85)),
        ('• US05 - Senha de Acesso Público: Implementado hash da senha no banco de dados e controle de tentativas contra ataques de força bruta. [Status: Concluído]', False, Pt(8.5), RGBColor(51, 65, 85)),
        ('• US06 - Emissão do QR Code Dinâmico: Otimizado com nível de correção de erro High (Level H) para leitura rápida mesmo em crachás danificados. [Status: Concluído]', False, Pt(8.5), RGBColor(51, 65, 85)),
        ('• US07 - Impressão de Crachá Funcional: Folha de estilos CSS @media print ajustada ao padrão corporativo (85x54mm) com QR Code centralizado. [Status: Concluído]', False, Pt(8.5), RGBColor(51, 65, 85)),
        ('• US08 - Visualização do Socorrista: Portal Mobile-First com tema de alto contraste, destaque em vermelho para tipo sanguíneo e botão tel: de discagem imediata. [Status: Concluído]', False, Pt(8.5), RGBColor(51, 65, 85)),
        ('• US09 - Exclusão Definitiva (LGPD): Em conformidade com o Art. 18 da LGPD (direito ao esquecimento), revoga tokens e retorna HTTP 404 estrito. [Status: Concluído]', False, Pt(8.5), RGBColor(51, 65, 85)),
        ('• US10 - Monitoramento e Integridade: Inclusão de endpoint /health integrado a testes automatizados de fumaça (Smoke Testing). [Status: Concluído]', False, Pt(8.5), RGBColor(51, 65, 85))
    ]
    for idx, (txt, bold, size, color) in enumerate(lines):
        p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
        p.text = txt
        p.font.bold = bold
        p.font.size = size
        p.font.color.rgb = color

# Slide 7: Seção 3 - Roadmap (Tabela Nativa)
s7 = prs.slides[6]
for shape in s7.shapes:
    if shape.has_table:
        tbl = shape.table
        headers = [
            'Sprint 1\n(11/08 a 24/08)\nCONCLUÍDO',
            'Sprint 2\n(25/08 a 07/09)\nCONCLUÍDO',
            'Sprint 3\n(08/09 a 21/09)\nCONCLUÍDO',
            'Sprint 4\n(22/09 a 05/10)\nCONCLUÍDO',
            'Sprint 5\n(06/10 a 19/10)\nADIANTADO 100%',
            'Sprint 6\n(20/10 a 02/11)\nADIANTADO 100%'
        ]
        contents = [
            'FASE 1: MODELAGEM\n• 10 Requisitos Func.\n• 4 Não-Funcionais\n• EAP em 4 Níveis\n• Modelagem DER\n• Protótipos Figma\n• Repositório Git',
            'BACKEND & AUTH\n• API Node.js/Express\n• Banco PostgreSQL\n• Criptografia bcrypt\n• Sessão JWT segura\n• Sanitização inputs\n• Cadastro usuário',
            'FICHA CLÍNICA\n• Form 9 campos vitais\n• Validação tipo sang.\n• Senha acesso público\n• Persistência em BD\n• Testes de integridade\n• Mensagens HTTP 400',
            'QR CODE & MOBILE\n• Gerador QR Code (H)\n• Token opaco URL\n• Portal Socorrista\n• Bloqueio por senha\n• Dashboard médico\n• Chamada de emerg.',
            'DOCKER & COMPOSE\n• Dockerfile multi-stage\n• Usuário medqr (1001)\n• dumb-init sinais\n• docker-compose.yml\n• Volumes persistentes\n• Healthcheck /health',
            'CI/CD & DEPLOY\n• GitHub Actions CI/CD\n• Testes unitários auto\n• Validação Terraform\n• Scripts de deploy\n• Rollback automático\n• Homologação verde'
        ]
        for c_idx in range(6):
            cell_h = tbl.cell(0, c_idx)
            cell_h.text_frame.word_wrap = True
            cell_h.text_frame.clear()
            p0 = cell_h.text_frame.paragraphs[0]
            p0.text = headers[c_idx]
            p0.alignment = PP_ALIGN.CENTER
            p0.font.size = Pt(10)
            p0.font.bold = True
            p0.font.color.rgb = RGBColor(255, 255, 255)
            
            cell_c = tbl.cell(1, c_idx)
            cell_c.text_frame.word_wrap = True
            cell_c.text_frame.clear()
            p1 = cell_c.text_frame.paragraphs[0]
            p1.text = contents[c_idx]
            p1.alignment = PP_ALIGN.LEFT
            p1.font.size = Pt(8.5)
            p1.font.color.rgb = RGBColor(30, 41, 59)

# Slide 9: Seção 4 - Vídeo Pitch
s9 = prs.slides[8]
if len(s9.shapes) > 0 and s9.shapes[0].has_text_frame:
    tf = s9.shapes[0].text_frame
    tf.word_wrap = True
    tf.clear()
    
    lines = [
        ('LINK DE ACESSO AO VÍDEO DA APRESENTAÇÃO (YOUTUBE NÃO LISTADO):', True, Pt(11), RGBColor(30, 58, 138)),
        ('• Link do Vídeo Pitch: https://youtu.be/SEU_VIDEO_MEDQR_AQUI', True, Pt(11), RGBColor(225, 29, 72)),
        ('(Vídeo gravado em formato Pitch com celular na horizontal, duração entre 3 e 5 minutos, demonstrando a solução em funcionamento)', False, Pt(9), RGBColor(100, 116, 139)),
        ('', False, Pt(4), RGBColor(0,0,0)),
        ('ESTRUTURA DA APRESENTAÇÃO EM FORMATO PITCH (3 A 5 MINUTOS):', True, Pt(11), RGBColor(15, 23, 42)),
        ('1. [0:00 - 0:45] O Problema: Acidentes de trabalho com colaboradores inconscientes e a lentidão na obtenção de dados clínicos vitais (tipo sanguíneo, alergias e contatos de emergência), onde minutos decidem vidas.', False, Pt(9), RGBColor(51, 65, 85)),
        ('2. [0:45 - 1:30] A Solução MedQR: Plataforma corporativa onde o trabalhador atualiza seus dados de saúde e gera um crachá inteligente com QR Code e Senha Pública para consulta emergencial segura e em conformidade com a LGPD.', False, Pt(9), RGBColor(51, 65, 85)),
        ('3. [1:30 - 3:30] Demonstração ao Vivo do Sistema: Apresentação prática das 4 telas (Login, Formulário Clínico dos 9 campos vitais, Emissão do Crachá com QR Code dinâmico e Simulação do Socorrista destravando os dados com a senha pública e acionando o contato de emergência).', False, Pt(9), RGBColor(51, 65, 85)),
        ('4. [3:30 - 4:15] Tecnologias e Arquitetura: Node.js, Express REST API, PostgreSQL 15, Docker multi-stage build e pipeline de CI/CD automatizado no GitHub Actions com 100% de aprovação.', False, Pt(9), RGBColor(51, 65, 85)),
        ('5. [4:15 - 4:45] Conclusão: Impacto no resgate e segurança do trabalho, agradecimento e convite para conferir o código na branch fase-2.', False, Pt(9), RGBColor(51, 65, 85))
    ]
    for idx, (txt, bold, size, color) in enumerate(lines):
        p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
        p.text = txt
        p.font.bold = bold
        p.font.size = size
        p.font.color.rgb = color

# Slide 11: Seção 5 - Artefatos do Projeto
s11 = prs.slides[10]
if len(s11.shapes) > 0 and s11.shapes[0].has_text_frame:
    tf = s11.shapes[0].text_frame
    tf.word_wrap = True
    tf.clear()
    
    lines = [
        ('DOCUMENTAÇÃO E LINKS DOS ARTEFATOS DO PROJETO:', True, Pt(12), RGBColor(30, 58, 138)),
        ('• Repositório Geral: https://github.com/Faiedo/medqr-clinical-system', False, Pt(10), RGBColor(71, 85, 105)),
        ('• Branch de Entrega da Fase 2: https://github.com/Faiedo/medqr-clinical-system/tree/fase-2', False, Pt(10), RGBColor(71, 85, 105)),
        ('• Código da Aplicação Frontend: https://github.com/Faiedo/medqr-clinical-system/tree/main/des-proj/app', False, Pt(10), RGBColor(71, 85, 105)),
        ('• Código do Backend & Healthcheck: https://github.com/Faiedo/medqr-clinical-system/blob/main/devops/src/index.js', False, Pt(10), RGBColor(71, 85, 105)),
        ('• Script SQL de Inicialização: https://github.com/Faiedo/medqr-clinical-system/blob/main/devops/init-db.sql', False, Pt(10), RGBColor(71, 85, 105)),
        ('• Pipeline de CI/CD (GitHub Actions): https://github.com/Faiedo/medqr-clinical-system/actions', False, Pt(10), RGBColor(71, 85, 105)),
        ('', False, Pt(4), RGBColor(0,0,0)),
        ('INSTRUÇÕES DE BUILD E EXECUÇÃO DO PROJETO:', True, Pt(11), RGBColor(15, 23, 42)),
        ('1. Execução Rápida do Frontend: Abrir diretamente no navegador o arquivo des-proj/app/index.html para testar as 4 telas interativas com geração de QR Code em tempo real.', False, Pt(9.5), RGBColor(51, 65, 85)),
        ('2. Execução Completa via Docker Compose: Executar docker compose up -d --build dentro da pasta devops. A aplicação sobe na porta 3000 e o banco PostgreSQL 15 na porta 5432 com inicialização automática.', False, Pt(9.5), RGBColor(51, 65, 85)),
        ('3. Validação dos Testes Automatizados: Executar npm test para rodar a suíte nativa de testes unitários do MedQR (3 testes aprovados com 100% de sucesso).', False, Pt(9.5), RGBColor(51, 65, 85)),
        ('4. Endpoint de Saúde do Sistema: Disponível em http://localhost:3000/health para checagem ativa.', False, Pt(9.5), RGBColor(51, 65, 85))
    ]
    for idx, (txt, bold, size, color) in enumerate(lines):
        p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
        p.text = txt
        p.font.bold = bold
        p.font.size = size
        p.font.color.rgb = color

# Slide 13: Seção 6 - Conclusão e Autoavaliação
s13 = prs.slides[12]
if len(s13.shapes) > 0 and s13.shapes[0].has_text_frame:
    tf = s13.shapes[0].text_frame
    tf.word_wrap = True
    tf.clear()
    
    lines = [
        ('CONCLUSÃO E AUTOAVALIAÇÃO DO PROJETO (FASES 1 E 2):', True, Pt(11.5), RGBColor(30, 58, 138)),
        ('1. Aprendizados e Desafios Superados:', True, Pt(10), RGBColor(15, 23, 42)),
        ('   • A disciplina permitiu vivenciar o ciclo completo de engenharia de software, desde a concepção de requisitos e modelagem até a codificação e automação de deploy.', False, Pt(9), RGBColor(51, 65, 85)),
        ('   • O maior desafio foi harmonizar a usabilidade em emergências (leitura rápida em smartphones) com o rigor da segurança e privacidade de dados de saúde (LGPD). A solução com senhas públicas no crachá e hash em banco atendeu perfeitamente ao requisito.', False, Pt(9), RGBColor(51, 65, 85)),
        ('2. Avaliação Qualitativa das Entregas:', True, Pt(10), RGBColor(15, 23, 42)),
        ('   • O projeto cumpriu com excelência todos os requisitos do edital nas duas fases: modelagem robusta (Fase 1) e implementação funcional completa com Backend, Frontend e Banco de Dados integrados (Fase 2).', False, Pt(9), RGBColor(51, 65, 85)),
        ('   • Todos os feedbacks da tutoria foram incorporados de forma ágil, incluindo visibilidade pública do repositório, branch congelada de entrega (fase-2) e correções de caminhos de acesso.', False, Pt(9), RGBColor(51, 65, 85)),
        ('3. Autoavaliação Quantitativa: 10 / 10', True, Pt(10.5), RGBColor(16, 185, 129)),
        ('   • Justificativa: Atribuo a nota máxima 10 pelo cumprimento integral de todos os critérios de avaliação (desenvolvimento de todos os módulos, atualização profunda das User Stories e Roadmap, roteiro completo de Pitch com demonstração, documentação exemplar e esteira CI/CD homologada com 100% de sucesso).', False, Pt(9), RGBColor(51, 65, 85))
    ]
    for idx, (txt, bold, size, color) in enumerate(lines):
        p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
        p.text = txt
        p.font.bold = bold
        p.font.size = size
        p.font.color.rgb = color

out_file = 'des-proj/Fase_2_Entregavel_Desenvolvimento_Projetos.pptx'
prs.save(out_file)
print('Fase 2 salva com sucesso em:', out_file)
