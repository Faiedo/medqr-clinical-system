import pptx, win32com.client, os

prs = pptx.Presentation('des-proj/Fase_2_Entregavel_Desenvolvimento_Projetos.pptx')

slides_to_remove = []
for i, slide in enumerate(prs.slides):
    is_guideline = False
    for shape in slide.slide_layout.shapes:
        if shape.has_text_frame and 'ORIENTAÇÕES' in shape.text_frame.text:
            is_guideline = True
            break
    for shape in slide.shapes:
        if shape.has_text_frame and 'ORIENTAÇÕES' in shape.text_frame.text:
            is_guideline = True
            break
    if i == len(prs.slides) - 1:
        has_content = any(s.has_text_frame and s.text_frame.text.strip() for s in slide.shapes) or any(s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE for s in slide.shapes)
        if not has_content:
            is_guideline = True

    if is_guideline:
        slides_to_remove.append(i)

print('Removendo slides de orientação:', slides_to_remove)
for idx in sorted(slides_to_remove, reverse=True):
    rId = prs.slides._sldIdLst[idx].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[idx]

out_pptx = 'des-proj/Fase_2_Entregavel_Desenvolvimento_Projetos_entrega.pptx'
prs.save(out_pptx)
print('Salvo arquivo de entrega limpo:', out_pptx)

# Convert to PDF via PowerPoint COM
ppt = win32com.client.Dispatch('PowerPoint.Application')
try:
    abs_in = os.path.abspath(out_pptx)
    abs_out = os.path.abspath('des-proj/Fase_2_Entregavel_Desenvolvimento_Projetos_entrega.pdf')
    deck = ppt.Presentations.Open(abs_in, True, False, False)
    deck.SaveAs(abs_out, 32)
    deck.Close()
    print('PDF da Fase 2 de Desenvolvimento de Projetos gerado com sucesso:', abs_out)
    print('Tamanho:', os.path.getsize(abs_out), 'bytes')
finally:
    ppt.Quit()
